from flask import Flask, request, jsonify, session, send_from_directory
from flask_cors import CORS
import bcrypt
import os
from datetime import datetime, date
from excel_db import (
    get_all_users,
    get_user_by_username,
    get_all_monitor_projects,
    get_projects_for_user,
    get_project_by_id,
    create_new_project_from_monitor,
    get_tasks,
    update_tasks_bulk,
    create_user,
    update_user,
    delete_user,
    get_master_projects,
    get_discipline_dirs,
    OWNER_MAP,
    update_monitor_cell,
    ROLE_TO_DEPT,      
    OWNER_MAP,
    update_monitor_timestamp,
    load_user_overdue_status,
)

app = Flask(__name__, static_folder=None)

# In-memory sheet cache (JSON sidecar handles persistence across restarts)
_sheet_cache = {}
# In-memory cache for master project lists (monitor file is read-only from outside)
_master_projects_cache = {}
app.secret_key = "scheduler_excel_secret_2024"
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_HTTPONLY"] = True

CORS(app, supports_credentials=True)

# Simple in-memory sheet cache
# ── Startup: build index ───────────────────────────────────────




FRONTEND_DIR = os.path.join(os.path.dirname(__file__), "..", "frontend")

# ── Serve frontend ─────────────────────────────────────────────
@app.route("/")
def index():
    return send_from_directory(FRONTEND_DIR, "index.html")

@app.route("/<path:filename>")
def static_files(filename):
    return send_from_directory(FRONTEND_DIR, filename)

# ── Auth helpers ───────────────────────────────────────────────
def get_current_user():
    username = session.get("username")
    if not username:
        return None
    return get_user_by_username(username)

def login_required(fn):
    from functools import wraps
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not get_current_user():
            return jsonify({"error": "Not authenticated"}), 401
        return fn(*args, **kwargs)
    return wrapper

# ── Auth routes ────────────────────────────────────────────────
@app.route("/api/auth/login", methods=["POST"])
def login():
    data = request.get_json()
    if not data or not data.get("username") or not data.get("password"):
        return jsonify({"error": "Username and password required"}), 400

    user = get_user_by_username(data["username"])
    if not user:
        return jsonify({"error": "Invalid credentials"}), 401

    if not bcrypt.checkpw(data["password"].encode(), user["password_hash"].encode()):
        return jsonify({"error": "Invalid credentials"}), 401

    session["username"] = user["username"]

    # Load overdue status for this user and store in session
    from excel_db import load_user_overdue_status
    session["overdue_map"] = load_user_overdue_status(user)

    return jsonify({
        "name":     user["name"],
        "username": user["username"],
        "role":     user["role"],
    })

@app.route("/api/auth/logout", methods=["POST"])
def logout():
    session.clear()
    return jsonify({"message": "Logged out"})

@app.route("/api/auth/me", methods=["GET"])
@login_required
def me():
    u = get_current_user()
    return jsonify({
        "name":       u["name"],
        "short_name": u["short_name"],
        "username":   u["username"],
        "role":       u["role"],
    })

# ── Projects routes ────────────────────────────────────────────
@app.route("/api/projects", methods=["GET"])
@login_required
def list_projects():
    user = get_current_user()
    projects = get_projects_for_user(user)  # role-aware internally
    return jsonify(projects)

@app.route("/api/projects/<project_id>", methods=["GET"])
@login_required
def get_project(project_id):
    user = get_current_user()
    project = get_project_by_id(project_id, role=user["role"])
    if not project:
        return jsonify({"error": "Project not found"}), 404

    # Access check — if project is in user's monitoring file they have access
    role = user["role"]
    if role != "admin" and not role.endswith("_head"):
        cache_key = f"all_master_projects_{role}"
        if cache_key not in _master_projects_cache:
            _master_projects_cache[cache_key] = get_all_monitor_projects(role)
        master_list = _master_projects_cache[cache_key]
        allowed_ids = {m.get("file_id") for m in master_list}
        if project_id not in allowed_ids:
            return jsonify({"error": "Access denied"}), 403

    return jsonify(project)



@app.route("/api/projects/<project_id>/tasks", methods=["GET"])
@login_required
def get_project_tasks(project_id):
    user = get_current_user()
    role = user["role"]
    owner_filter = OWNER_MAP.get(role)
    tasks = get_tasks(project_id, owner_filter=owner_filter, role=role)
    return jsonify(tasks)


@app.route("/api/projects/<project_id>/tasks", methods=["PUT"])
@login_required
def save_tasks(project_id):
    user = get_current_user()
    role = user["role"]
    data = request.get_json()

    if not data or not isinstance(data, list):
        return jsonify({"error": "Expected list of task updates"}), 400

    # Verify each task's owner matches user role (unless admin/head/pm)
    if role not in ("admin", "head", "pm"):
        allowed_owner = OWNER_MAP.get(role)
        for item in data:
            if "_row" in item:
                continue
            if item.get("owner") != allowed_owner:
                return jsonify({"error": f"You can only update {allowed_owner} tasks"}), 403

    ok, msg, now_str = update_tasks_bulk(project_id, data, role=role)
    if not ok:
        return jsonify({"error": msg}), 404

    # AFTER SAVE: Refresh the overdue_map for this project in session
    from excel_db import read_sysmemory_json
    overdue_map = session.get("overdue_map", {})
    
    # Refresh this project's data
    fresh_task_map = read_sysmemory_json(project_id, role)
    if fresh_task_map:
        overdue_map[project_id] = fresh_task_map
    else:
        # If no data, remove from map
        overdue_map.pop(project_id, None)
    
    session["overdue_map"] = overdue_map

    # Reload the in-memory cache
    from excel_db import _read_sheet_cache, get_discipline_dirs
    _, _, sched_cache, _ = get_discipline_dirs(role)
    fresh = _read_sheet_cache(project_id, sched_cache)

    sheet_cache_key = f"{role}:{project_id}" 

    if fresh:
        _sheet_cache[sheet_cache_key] = fresh
    elif sheet_cache_key in _sheet_cache:
        if now_str:
            _sheet_cache[sheet_cache_key]["last_modified"] = now_str

    return jsonify({"message": msg, "last_modified": now_str})

@app.route("/api/projects/<project_id>/sheet", methods=["GET"])
@login_required
def get_sheet_data(project_id):
    """Return raw cell data for the Excel-mirror UI. Cached in memory."""
    from excel_db import get_raw_sheet, get_discipline_dirs
    user = get_current_user()
    role = user["role"]
    is_readonly = role in ("admin", "head")

    # NEW: role-aware in-memory key
    sheet_cache_key = f"{role}:{project_id}"

    # Check in-memory cache first (fastest)
    if sheet_cache_key in _sheet_cache:
        data = _sheet_cache[sheet_cache_key]
        if is_readonly:
            # Return a copy with editable flags stripped — don't mutate the cache
            import copy
            data = copy.deepcopy(data)
            for cell in data.get("cells", {}).values():
                cell.pop("editable", None)
        return jsonify(data)

    # CHANGED: capture sched_cache too
    projects_dir, _, sched_cache, _ = get_discipline_dirs(role)
    
    fpath = os.path.join(projects_dir, project_id + ".xlsx")
    if not os.path.exists(fpath):
        fpath_b = os.path.join(projects_dir, project_id + ".xlsb")
        if os.path.exists(fpath_b):
            fpath = fpath_b
        else:
            return jsonify({"error": "Project file not found"}), 404
    try:
        data = get_raw_sheet(fpath, role=role, sched_cache=sched_cache)
        # Inject last_modified from file mtime if get_raw_sheet didn't already
        if not data.get("last_modified"):
            from datetime import datetime as _dt
            try:
                ts = os.path.getmtime(fpath)
                data["last_modified"] = _dt.fromtimestamp(ts).strftime("%d %b %Y, %I:%M %p")
            except Exception:
                data["last_modified"] = None
        _sheet_cache[sheet_cache_key] = data  # store in memory
        if is_readonly:
            import copy
            data = copy.deepcopy(data)
            for cell in data.get("cells", {}).values():
                cell.pop("editable", None)
        return jsonify(data)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route("/api/master/projects", methods=["GET"])
@login_required
def master_projects():
    user = get_current_user()
    role = user["role"]
    cache_key = f"all_master_projects_{role}"   # <-- role-specific key
    if cache_key not in _master_projects_cache:
        from excel_db import get_all_monitor_projects
        _master_projects_cache[cache_key] = get_all_monitor_projects(role)
    return jsonify(_master_projects_cache[cache_key])

@app.route("/api/monitor/sheet", methods=["GET"])
@login_required
def get_monitor_sheet():
    """Return the raw sheet data for the department monitoring file."""
    from excel_db import get_monitor_sheet_data, find_monitor_file, load_user_overdue_status
    import os
    
    user = get_current_user()
    role = user.get("role", "sw_tl")
    
    # Map role to department
    department = ROLE_TO_DEPT.get(role, "SW")
    
    # Use find_monitor_file instead
    fpath = find_monitor_file(department)
    
    print(f"[Monitor] Looking for: {fpath}")
    
    if not os.path.exists(fpath):
        print(f"[Monitor] File not found: {fpath}")
        return jsonify({"error": f"No monitoring file found for {department}"}), 404
    
    try:
        data = get_monitor_sheet_data(fpath, department)
        
        # ALWAYS refresh overdue_map when monitor loads
        fresh_overdue_map = load_user_overdue_status(user)
        session["overdue_map"] = fresh_overdue_map
        
        return jsonify({
            "sheet": data,
            "overdue": fresh_overdue_map
        })
    except Exception as e:
        print(f"[Monitor] Error: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500
        
@app.route("/api/projects/<project_id>/monitor-timestamp", methods=["POST"])
@login_required
def save_monitor_timestamp(project_id):
    """Write the save timestamp into the BA column of the monitor JSON/Excel."""
    user = get_current_user()
    role = user["role"]

    data = request.get_json()
    display_ts = (data or {}).get("timestamp", "")

    # Convert display format "29 Mar 2026, 08:04 PM" -> raw "29-03-2026 20:04:00"
    # to match the Excel cell format dd-mm-yyyy hh:mm:ss
    raw_ts = display_ts
    try:
        raw_ts = datetime.strptime(display_ts, "%d %b %Y, %I:%M %p").strftime("%d-%m-%Y %H:%M:%S")
    except Exception:
        raw_ts = datetime.now().strftime("%d-%m-%Y %H:%M:%S")

    ok = update_monitor_timestamp(project_id, raw_ts, role=role)
    if not ok:
        # Non-fatal — project may not be in this user's monitor file
        return jsonify({"message": "Project not found in monitor, skipped"}), 200

    # Patch the in-memory master projects cache in-place so the sidebar
    # reflects the updated stale status immediately — no disk read needed.
    # With:
    cache_key = f"all_master_projects_{role}"
    cached = _master_projects_cache.get(cache_key, [])
    for entry in cached:
        if entry.get("file_id") == project_id:
            entry["stale"] = False
            break

    return jsonify({"message": "Monitor timestamp updated", "timestamp": raw_ts})


@app.route("/api/debug", methods=["GET"])
@login_required
def debug():
    user = get_current_user()
    from excel_db import get_all_projects
    projects = get_all_projects()
    return jsonify({
        "logged_in_user": {
            "name":       user["name"],
            "username":   user["username"],
            "short_name": user.get("short_name", "MISSING"),
            "role":       user["role"],
        },
        "projects": [{
            "id":           p["id"],
            "pm_name":      p.get("pm_name", ""),
            "hw_tl_name":   p.get("hw_tl_name", ""),
            "sw_tl_name":   p.get("sw_tl_name", ""),
            "mfg_tl_name":  p.get("mfg_tl_name", ""),
        } for p in projects]
    })


@app.route("/api/admin/rebuild-index", methods=["POST"])
@login_required
def rebuild_index():
    user = get_current_user()
    if user["role"] != "admin":
        return jsonify({"error": "Access denied"}), 403
    # Clear in-memory sheet cache so all projects reload fresh
    _sheet_cache.clear()
    _master_projects_cache.clear()
    return jsonify({"message": "Cache cleared successfully"})

@app.route("/api/admin/clear-monitor-cache", methods=["POST"])
@login_required
def clear_monitor_cache():
    user = get_current_user()
    if user["role"] != "admin":
        return jsonify({"error": "Access denied"}), 403
    
    # Clear in-memory monitor cache
    global _sheet_cache
    keys_to_remove = [k for k in _sheet_cache.keys() if k.startswith("monitor_")]
    for key in keys_to_remove:
        del _sheet_cache[key]
    
    # Also delete JSON cache files for monitor
    from excel_db import DATA_DIR
    import os
    
    # Delete monitor cache for all departments
    departments = ["SW", "HW", "MFG", "PM"]
    for dept in departments:
        from excel_db import get_cache_path
        monitor_cache_dir = get_cache_path(dept, "Monitoring")
        if os.path.exists(monitor_cache_dir):
            for filename in os.listdir(monitor_cache_dir):
                if filename.endswith(".json"):
                    filepath = os.path.join(monitor_cache_dir, filename)
                    try:
                        os.remove(filepath)
                        print(f"Deleted monitor cache: {filepath}")
                    except Exception as e:
                        print(f"Could not delete {filepath}: {e}")
    
    return jsonify({"message": "Monitor cache cleared successfully"})

# __NEW PROJECT CREATION LOGIC__________________________________

@app.route('/api/monitor/create-project', methods=['POST'])
def create_project_from_monitor():
    """Create new project from monitor UI (PM Head only)"""
    user = get_current_user()
    if not user or user.get('role') != 'pm_head':
        return jsonify({'error': 'Unauthorized'}), 401
    
    data = request.json
    or_number = data.get('or_number')
    section = data.get('section')
    ov_value = data.get('ov_value')
    assign_to = data.get('assign_to')
    
    if not all([or_number, section, ov_value, assign_to]):
        return jsonify({'error': 'Missing required fields'}), 400
    
    # Call excel_db function to create project
    from excel_db import create_new_project_from_monitor
    result = create_new_project_from_monitor(or_number, section, ov_value, assign_to, user)
    
    return jsonify(result)

@app.route("/api/projects/<project_id>/setup", methods=["GET", "POST"])
@login_required
def project_setup(project_id):
    """Get setup fields or save setup data for a new project"""
    user = get_current_user()
    
    # Only PM or pm_head can access setup
    if user["role"] not in ("pm", "pm_head"):
        return jsonify({"error": "Access denied"}), 403
    
    from excel_db import get_setup_fields, write_project_setup_data, get_discipline_dirs, _read_sheet_cache
    import os
    
    # Get department for this user
    from excel_db import ROLE_TO_DEPT
    dept = ROLE_TO_DEPT.get(user["role"], "PM")
    
    # Get the schedule cache path to check if JSON exists
    _, _, sched_cache, _ = get_discipline_dirs(user["role"])
    cache_path = os.path.join(sched_cache, project_id + "_sheet.json")
    
    # If GET request: return the setup form configuration
    if request.method == "GET":
        # If JSON cache already exists, project is already set up
        if os.path.exists(cache_path):
            return jsonify({
                "already_setup": True,
                "message": "Project already has a cache file"
            })
        
        # Get the setup fields configuration
        setup_fields = get_setup_fields(dept)
        
        # Get task names for rows that have S, U, or AD fields
        from openpyxl import load_workbook
        projects_dir, _, _, _ = get_discipline_dirs(user["role"])
        file_path = os.path.join(projects_dir, project_id + ".xlsx")
        
        if not os.path.exists(file_path):
            file_path = os.path.join(projects_dir, project_id + ".xlsb")
        
        task_names = {}
        if os.path.exists(file_path):
            wb = load_workbook(file_path, data_only=True)
            ws = wb.active
            # Read task names from column I for rows 9-55
            for row in range(9, 56):
                task_name = ws[f"I{row}"].value
                if task_name:
                    task_names[row] = str(task_name).strip()
            wb.close()
        
        return jsonify({
            "already_setup": False,
            "fields": setup_fields,
            "task_names": task_names,
            "project_id": project_id
        })
    
    # If POST request: save the setup data
    if request.method == "POST":
        form_data = request.get_json()
        
        if not form_data:
            return jsonify({"error": "No data provided"}), 400
        
        # Write data to Excel and generate cache
        success, message, last_modified = write_project_setup_data(project_id, form_data, user["role"])
        
        if not success:
            return jsonify({"error": message}), 500
        
        # Also update monitor timestamp
        if last_modified:
            from excel_db import update_monitor_timestamp
            update_monitor_timestamp(project_id, last_modified, user["role"])
        
        return jsonify({
            "success": True,
            "message": message,
            "last_modified": last_modified
        })

# ── User management (admin only) ──────────────────────────────
VALID_ROLES = ["admin", "sw_head", "hw_head", "mfg_head", "pm_head", "pm", "hw_tl", "sw_tl", "mfg_tl"]

@app.route("/api/users", methods=["GET"])
@login_required
def list_users():
    user = get_current_user()
    if user["role"] != "admin":
        return jsonify({"error": "Access denied"}), 403
    users = get_all_users()
    # Don't send password hashes to frontend
    return jsonify([{
        "name":       u["name"],
        "short_name": u["short_name"],
        "username":   u["username"],
        "role":       u["role"],
    } for u in users])

@app.route("/api/users", methods=["POST"])
@login_required
def add_user():
    user = get_current_user()
    if user["role"] != "admin":
        return jsonify({"error": "Access denied"}), 403

    data = request.get_json()
    for f in ["name", "short_name", "username", "password", "role"]:
        if not data.get(f):
            return jsonify({"error": f"{f} is required"}), 400
    if data["role"] not in VALID_ROLES:
        return jsonify({"error": "Invalid role"}), 400

    ok, msg = create_user(data["name"], data["short_name"], data["username"], data["password"], data["role"])
    if not ok:
        return jsonify({"error": msg}), 409
    return jsonify({"message": msg}), 201

@app.route("/api/users/<username>", methods=["PUT"])
@login_required
def edit_user(username):
    user = get_current_user()
    if user["role"] != "admin":
        return jsonify({"error": "Access denied"}), 403

    data = request.get_json()
    if "role" in data and data["role"] not in VALID_ROLES:
        return jsonify({"error": "Invalid role"}), 400

    ok, msg = update_user(username, data)
    if not ok:
        return jsonify({"error": msg}), 404
    return jsonify({"message": msg})

@app.route("/api/users/<username>", methods=["DELETE"])
@login_required
def remove_user(username):
    user = get_current_user()
    if user["role"] != "admin":
        return jsonify({"error": "Access denied"}), 403
    if username == user["username"]:
        return jsonify({"error": "Cannot delete your own account"}), 400

    ok, msg = delete_user(username)
    if not ok:
        return jsonify({"error": msg}), 404
    return jsonify({"message": msg})

@app.route("/api/refresh-file-status", methods=["POST"])
def refresh_file_status():
    from excel_db import get_all_monitor_projects, find_schedule_folder, DEPT_CONFIG
    import os
    
    user = get_current_user()
    if not user:
        return jsonify({"error": "Not authenticated"}), 401
    
    role = user.get("role", "sw_tl")
    projects = get_all_monitor_projects(role)
    
    # Map role to department
    
    dept = ROLE_TO_DEPT.get(role, "SW")
    projects_dir = find_schedule_folder(dept)
    
    # Get the correct prefix for this department
    prefix = DEPT_CONFIG[dept]["sched_prefix"]
    
    print("=== REFRESH DEBUG ===")
    
    for project in projects:
        project_id = project.get("project_id")
        if project_id:
            # Use department-specific prefix instead of hardcoded "SWESch_"
            normalized_id = prefix + "_" + project_id.replace("/", "_")
            excel_path_xlsx = os.path.join(projects_dir, normalized_id + ".xlsx")
            excel_path_xlsb = os.path.join(projects_dir, normalized_id + ".xlsb")
            file_exists = os.path.exists(excel_path_xlsx) or os.path.exists(excel_path_xlsb)
            
            print(f"Project ID: {project_id}")
            print(f"  Normalized: {normalized_id}")
            print(f"  Looking for: {excel_path_xlsx}")
            print(f"  Exists: {file_exists}")
            
            project["file_exists"] = file_exists
    
    return jsonify({"projects": projects})

@app.route("/api/monitor/cell", methods=["POST"])
@login_required
def update_monitor_cell():
    """Save a single cell edit in the monitor file"""
    user = get_current_user()
    role = user["role"]
    
    data = request.get_json()
    if not data:
        return jsonify({"error": "No data provided"}), 400
    
    col = data.get("col")
    row = data.get("row")
    value = data.get("value")
    
    if not col or not row:
        return jsonify({"error": "Column and row required"}), 400
    
    department = ROLE_TO_DEPT.get(role, "SW")
    
    # Update monitor cache and queue Excel write
    coord = f"{col}{row}"
    
    try:
        from excel_db import update_monitor_cell
        result = update_monitor_cell(department, coord, value, role)
        if result:
            return jsonify({"message": "Saved", "coord": coord, "value": value})
        else:
            return jsonify({"error": "Failed to save"}), 500
    except Exception as e:
        print(f"Error saving monitor cell: {e}")
        return jsonify({"error": str(e)}), 500

if __name__ == "__main__":
    print("Starting Project Scheduler on http://localhost:5000")
    app.run(debug=True, host="0.0.0.0", port=5000)
